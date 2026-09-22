"""Office tool — offline docx/xlsx read+write (openpyxl / python-docx), the
deterministic file-deliverable channel: gateway chats deliver with
send_file, desktop lands them in the workspace (relative paths are
workspace-resolved like write_file's)."""

import importlib.util
import logging
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from tools import registry, tool_error, tool_result

logger = logging.getLogger(__name__)

_ACTIONS = ("excel_write", "excel_read", "word_write", "word_read")

# Column auto-width bounds for excel_write.
_COL_MIN, _COL_MAX = 8, 48


def _check() -> bool:
    return (importlib.util.find_spec("openpyxl") is not None
            or importlib.util.find_spec("docx") is not None)


def _resolve(path: str, kw: dict) -> Path:
    from core.support.paths import resolve_path
    return Path(str(resolve_path(path, kw.get("parent_agent"))))


def _cell_out(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, Decimal):
        return float(v)
    return v


def _display_width(s: str) -> int:
    # CJK glyphs count double toward column width.
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in s)


def _excel_write(args: dict, kw: dict) -> str:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return tool_error("openpyxl is not installed — excel unavailable in this environment")

    path = str(args.get("path") or "").strip()
    if not path:
        return tool_error("path is required (e.g. report.xlsx; relative = workspace)")
    sheets = args.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        return tool_error("sheets must be a non-empty array")

    out = _resolve(path, kw)
    if out.suffix.lower() != ".xlsx":
        out = out.with_suffix(".xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)
    for i, sh in enumerate(sheets):
        name = str((sh or {}).get("name") or f"Sheet{i + 1}")
        rows = (sh or {}).get("rows")
        if not isinstance(rows, list):
            return tool_error(f"sheet '{name}': rows must be an array of arrays")
        ws = wb.create_sheet(title=name[:31])
        for row in rows:
            if not isinstance(row, list):
                return tool_error(f"sheet '{name}': each row must be an array")
            ws.append([_cell_out(v) for v in row])
        if (sh or {}).get("header") and rows:
            for c in ws[1]:
                c.font = Font(bold=True)
            ws.freeze_panes = "A2"
        for col in ws.columns:
            width = max((_display_width(str(c.value or "")) for c in col), default=0)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, _COL_MIN), _COL_MAX)
    wb.save(out)

    logger.info("office [excel_write] -> %s", out)
    return tool_result(success=True, action="excel_write", path=str(out),
                       sheets=len(sheets))


def _excel_read(args: dict, kw: dict) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return tool_error("openpyxl is not installed — excel unavailable in this environment")

    path = str(args.get("path") or "").strip()
    if not path:
        return tool_error("path is required")
    p = _resolve(path, kw)
    if not p.exists():
        return tool_error(f"File not found: {p}")

    try:
        max_rows = int(args.get("max_rows") or 200)
    except (TypeError, ValueError):
        max_rows = 200
    max_rows = max(1, min(max_rows, 5000))

    wb = load_workbook(p, data_only=True, read_only=True)
    sheet = args.get("sheet")
    if sheet is None:
        ws = wb.active
    elif isinstance(sheet, int):
        ws = wb.worksheets[sheet] if -len(wb.worksheets) <= sheet < len(wb.worksheets) else None
    else:
        ws = wb[sheet] if str(sheet) in wb.sheetnames else None
    if ws is None:
        return tool_error(f"sheet not found: {sheet} (available: {wb.sheetnames})")

    rows, total = [], 0
    for row in ws.iter_rows(values_only=True):
        total += 1
        if len(rows) < max_rows:
            rows.append([_cell_out(v) for v in row])
    wb.close()

    return tool_result(success=True, action="excel_read", path=str(p),
                       sheet=ws.title, total_rows=total, returned_rows=len(rows),
                       truncated=total > len(rows), rows=rows)


def _word_write(args: dict, kw: dict) -> str:
    try:
        from docx import Document
    except ImportError:
        return tool_error("python-docx is not installed — word unavailable in this environment")

    path = str(args.get("path") or "").strip()
    if not path:
        return tool_error("path is required (e.g. report.docx; relative = workspace)")
    blocks = args.get("blocks")
    if not isinstance(blocks, list) or not blocks:
        return tool_error("blocks must be a non-empty array")

    out = _resolve(path, kw)
    if out.suffix.lower() != ".docx":
        out = out.with_suffix(".docx")
    out.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    for b in blocks:
        b = b or {}
        kind = str(b.get("type") or "paragraph")
        if kind == "heading":
            try:
                level = max(1, min(int(b.get("level") or 1), 4))
            except (TypeError, ValueError):
                level = 1
            doc.add_heading(str(b.get("text") or ""), level=level)
        elif kind == "bullet":
            for item in ([b.get("text")] if b.get("text") else b.get("items") or []):
                doc.add_paragraph(str(item), style="List Bullet")
        elif kind == "table":
            rows = b.get("rows")
            if not isinstance(rows, list) or not rows:
                return tool_error("table block requires rows (array of arrays)")
            t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            t.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, val in enumerate(row):
                    cell = t.cell(ri, ci)
                    cell.text = str(val)
                    if ri == 0 and b.get("header", True):
                        for run in cell.paragraphs[0].runs or [cell.paragraphs[0].add_run("")]:
                            run.font.bold = True
        else:
            doc.add_paragraph(str(b.get("text") or ""))

    doc.save(out)
    logger.info("office [word_write] -> %s", out)
    return tool_result(success=True, action="word_write", path=str(out),
                       blocks=len(blocks))


def _word_read(args: dict, kw: dict) -> str:
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        return tool_error("python-docx is not installed — word unavailable in this environment")

    path = str(args.get("path") or "").strip()
    if not path:
        return tool_error("path is required")
    p = _resolve(path, kw)
    if not p.exists():
        return tool_error(f"File not found: {p}")

    try:
        max_chars = int(args.get("max_chars") or 20000)
    except (TypeError, ValueError):
        max_chars = 20000
    max_chars = max(500, min(max_chars, 100000))

    doc = Document(str(p))
    lines = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            if not para.text.strip():
                continue
            if para.style.name.startswith("Heading"):
                try:
                    level = int(para.style.name.split()[-1])
                except ValueError:
                    level = 2
                lines.append("#" * max(1, min(level, 4)) + " " + para.text)
            elif para.style.name == "List Bullet":
                lines.append("- " + para.text)
            else:
                lines.append(para.text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, doc)
            for row in table.rows:
                lines.append("| " + " | ".join(c.text.strip() for c in row.cells) + " |")
    text = "\n".join(lines)
    truncated = len(text) > max_chars
    return tool_result(success=True, action="word_read", path=str(p),
                       text=text[:max_chars], truncated=truncated)


def _office(args: dict, **kw) -> str:
    action = str(args.get("action") or "").lower()
    if action not in _ACTIONS:
        return tool_error(f"action must be one of: {', '.join(_ACTIONS)}")
    if action == "excel_write":
        return _excel_write(args, kw)
    if action == "excel_read":
        return _excel_read(args, kw)
    if action == "word_write":
        return _word_write(args, kw)
    return _word_read(args, kw)


registry.register(
    name="office",
    schema={
        "type": "function",
        "function": {
            "name": "office",
            "description": (
                "Create and read office documents offline (xlsx via openpyxl, "
                "docx via python-docx) — the deterministic way to deliver "
                "real file artifacts (reports, data exports, memos). Gateway "
                "chats deliver the file with send_file(file_path=...); on "
                "desktop a relative path lands in the conversation's "
                "workspace. Four actions: excel_write (sheets of rows, "
                "optional header styling + column autofit), excel_read "
                "(rows back as JSON), word_write (heading/paragraph/bullet/"
                "table blocks), word_read (markdown-ish text back)."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "action": {"type": "string", "enum": list(_ACTIONS)},
                    "path": {"type": "string",
                             "description": "File path (.xlsx/.docx); relative = workspace"},
                    "sheets": {
                        "type": "array",
                        "description": "excel_write: [{name, rows: [[...]], header?: bool}]",
                        "items": {"type": "object"},
                    },
                    "sheet": {"type": "string",
                              "description": "excel_read: sheet name or index (default active)"},
                    "max_rows": {"type": "integer",
                                 "description": "excel_read row cap (default 200)"},
                    "blocks": {
                        "type": "array",
                        "description": "word_write: [{type: heading|paragraph|bullet|table, "
                                       "text, level, items, rows, header}]",
                        "items": {"type": "object"},
                    },
                    "max_chars": {"type": "integer",
                                  "description": "word_read text cap (default 20000)"},
                },
                "required": ["action"],
            },
        },
    },
    handler=lambda args, **kw: _office(args, **kw),
    check_fn=_check,
    toolset="office",
)
